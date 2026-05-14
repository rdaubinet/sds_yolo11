"""
Convert Markdown report to Word document format.
"""

import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def parse_markdown_table(lines):
    """Parse markdown table and return data."""
    if not lines or len(lines) < 2:
        return None
    
    # Extract headers (first line)
    headers = [cell.strip() for cell in lines[0].split('|') if cell.strip()]
    
    # Skip separator line (second line)
    # Extract data rows (remaining lines)
    data_rows = []
    for line in lines[2:]:
        if not line.strip():
            break
        cells = [cell.strip() for cell in line.split('|') if cell.strip()]
        if cells:
            data_rows.append(cells)
    
    return headers, data_rows

def add_formatted_text(paragraph, text):
    """Add text with markdown formatting (bold, italic, code)."""
    # Pattern to match **bold**, *italic*, `code`
    pattern = r'(\*\*.*?\*\*|\*.*?\*|`.*?`)'
    parts = re.split(pattern, text)
    
    for part in parts:
        if not part:
            continue
        
        if part.startswith('**') and part.endswith('**'):
            # Bold text
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('`') and part.endswith('`'):
            # Code text
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
        elif part.startswith('*') and part.endswith('*'):
            # Italic text
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            # Normal text
            paragraph.add_run(part)

def convert_markdown_to_docx(md_file, docx_file):
    """Convert markdown file to Word document."""
    
    # Create document
    doc = Document()
    
    # Set document margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    
    # Read markdown file
    with open(md_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    i = 0
    in_table = False
    table_lines = []
    in_code_block = False
    code_lines = []
    
    while i < len(lines):
        line = lines[i].rstrip()
        
        # Handle code blocks
        if line.startswith('```'):
            if in_code_block:
                # End code block
                if code_lines:
                    p = doc.add_paragraph('\n'.join(code_lines))
                    p.style = 'Intense Quote'
                    for run in p.runs:
                        run.font.name = 'Courier New'
                        run.font.size = Pt(9)
                code_lines = []
                in_code_block = False
            else:
                # Start code block
                in_code_block = True
            i += 1
            continue
        
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
        
        # Handle tables
        if line.startswith('|') and '|' in line[1:]:
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(line)
            i += 1
            continue
        else:
            if in_table:
                # End of table, process it
                result = parse_markdown_table(table_lines)
                if result:
                    headers, data_rows = result
                    
                    # Create table in document
                    table = doc.add_table(rows=1 + len(data_rows), cols=len(headers))
                    table.style = 'Light Grid Accent 1'
                    
                    # Add headers
                    hdr_cells = table.rows[0].cells
                    for idx, header in enumerate(headers):
                        hdr_cells[idx].text = header
                        # Bold header
                        for paragraph in hdr_cells[idx].paragraphs:
                            for run in paragraph.runs:
                                run.bold = True
                    
                    # Add data rows
                    for row_idx, row_data in enumerate(data_rows):
                        row_cells = table.rows[row_idx + 1].cells
                        for col_idx, cell_data in enumerate(row_data):
                            if col_idx < len(row_cells):
                                row_cells[col_idx].text = cell_data
                    
                    # Add spacing after table
                    doc.add_paragraph()
                
                in_table = False
                table_lines = []
        
        # Handle headings
        if line.startswith('#'):
            heading_match = re.match(r'^(#{1,6})\s+(.+)$', line)
            if heading_match:
                level = len(heading_match.group(1))
                title = heading_match.group(2)
                
                if level == 1:
                    heading = doc.add_heading(title, level=0)
                elif level <= 4:
                    heading = doc.add_heading(title, level=level-1)
                else:
                    # Use paragraph for deeper levels
                    p = doc.add_paragraph()
                    run = p.add_run(title)
                    run.bold = True
                    run.font.size = Pt(11)
                
                i += 1
                continue
        
        # Handle horizontal rules
        if line.strip() in ['---', '***', '___']:
            doc.add_paragraph('_' * 80)
            i += 1
            continue
        
        # Handle unordered lists
        if line.startswith('- ') or line.startswith('* '):
            text = line[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            add_formatted_text(p, text)
            i += 1
            continue
        
        # Handle ordered lists
        if re.match(r'^\d+\.\s+', line):
            text = re.sub(r'^\d+\.\s+', '', line).strip()
            p = doc.add_paragraph(style='List Number')
            add_formatted_text(p, text)
            i += 1
            continue
        
        # Handle empty lines
        if not line.strip():
            # Only add paragraph break if previous wasn't empty
            if i > 0 and lines[i-1].strip():
                doc.add_paragraph()
            i += 1
            continue
        
        # Handle regular paragraphs
        if line.strip():
            p = doc.add_paragraph()
            add_formatted_text(p, line)
        
        i += 1
    
    # Save document
    doc.save(docx_file)
    print(f"Successfully converted {md_file} to {docx_file}")

if __name__ == "__main__":
    md_file = "runs/detect/runs/detect/test_batch_best/PERFORMANCE_COMPARISON_REPORT.md"
    docx_file = "runs/detect/runs/detect/test_batch_best/PERFORMANCE_COMPARISON_REPORT.docx"
    
    convert_markdown_to_docx(md_file, docx_file)
